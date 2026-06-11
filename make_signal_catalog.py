# -*- coding: utf-8 -*-
"""
Draft-token selection signal catalog -> Excel.
다중 proposer(EAGLE3 / suffix / extension) 환경에서 budget 내 토큰 선정을 위한
'후보 신호'를 전수 정리한 카탈로그. (알고리즘 설계 전, 후보 수집 단계)
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# columns: No, Category, Signal, Proposer, Granularity, PyType, Range, Base/Derived,
#          Availability, CodeRef, Example, Description(KR)
HDR = ["No", "Category", "Signal (English)", "Proposer", "Granularity",
       "Python Type", "Range / Values", "Base/Derived", "Availability",
       "Code Reference", "Example", "Description (KR)"]

# Availability legend:
#  AVAIL  = 지금 바로 사용 가능 (capture/Draft에 이미 노출)
#  PLUMB  = native지만 plumbing 필요 (C++ Node / cache.py 내부에 존재, 미추출)
#  HIDDEN = 모델 내부에 존재하나 인터페이스 미노출 (deep patch 필요)
#  NONE   = 존재하지 않음 (확인 완료)

rows = [
# ---------------------------------------------------------------- A. 확률 / Confidence
["A. Confidence (확률)", "path_draft_p_t", "EAGLE3", "per-node",
 "list[float]", "[0.0, 1.0]", "Base",
 "AVAIL", "oracle_patch.py:301",
 "[0.92, 0.81, 0.74, ...]",
 "EAGLE3 드래프트 모델이 verify 전에 내부 ranking으로 계산한 root→node 누적 확률(경로상 per-edge 조건부 확률의 곱). 한 노드가 실제로 채택될 누적 신뢰도. inference-time에 사용 가능한 EAGLE3의 사실상 유일한 native 확률 신호."],

["A. Confidence (확률)", "eagle3_edge_prob (= path_p_t[i]/path_p_t[parent])", "EAGLE3", "per-edge",
 "float", "[0.0, 1.0]", "Derived(자명)", "AVAIL", "run_tree_oracle_sim.py:1567",
 "0.88",
 "부모 대비 한 edge의 조건부 확률. path_draft_p_t를 부모값으로 나눠 자명하게 복원. 누적값과 달리 '이 한 스텝'의 confidence만 분리한 값."],

["A. Confidence (확률)", "probs (suffix cumulative)", "suffix", "per-node",
 "list[float]", "[0.0, 1.0]", "Native-derived", "AVAIL", "suffix_tree.h:112 / Draft.probs",
 "[0.5, 0.33, 0.3, ...]",
 "Suffix tree에서 root→node 누적 확률. 각 edge가 child.count/parent.count(빈도 비율)이고 그 곱. 즉 빈도 기반 경로 확률. 경로를 따라 단조 감소."],

["A. Confidence (확률)", "suffix_edge_prob (= child.count/parent.count)", "suffix", "per-edge",
 "float", "[0.0, 1.0]", "Base", "AVAIL", "suffix_tree.h:117 / Draft.counts",
 "0.66",
 "Suffix tree에서 한 edge의 local 조건부 확률. counts로부터 직접 복원(child count / parent count). 누적 probs보다 '이 분기 자체'의 빈도 우세도를 직접 표현."],

# ---------------------------------------------------------------- B. 분포 형태 / Distributional shape
["B. Distributional shape (분포 형태)", "eagle3_draft_entropy", "EAGLE3", "per-node(parent 분포)",
 "float", "[0.0, log(V)]", "Derived(native logit)", "PLUMB", "oracle_patch.py:362 (flat_scores)",
 "2.7",
 "EAGLE3 드래프트 분포의 Shannon entropy. flat_scores(전체 후보 로짓)가 native하게 존재하나 현재 clone/저장 안 함. 낮으면 모델이 다음 토큰을 확신, 높으면 모호. top-1 확률 하나보다 정보량이 큼."],

["B. Distributional shape (분포 형태)", "eagle3_top1_top2_margin", "EAGLE3", "per-node(parent 분포)",
 "float", "[0.0, 1.0]", "Derived(native logit)", "PLUMB", "oracle_patch.py:362 (flat_scores)",
 "0.41",
 "1등과 2등 후보 토큰 확률의 차이(gap). margin이 크면 1등 토큰이 압도적으로 확신, 작으면 경합. flat_scores에서 topk(2)로 즉시 유도. 현재 미추출."],

["B. Distributional shape (분포 형태)", "eagle3_topk_mass", "EAGLE3", "per-node(parent 분포)",
 "float", "[0.0, 1.0]", "Derived(native logit)", "PLUMB", "oracle_patch.py:362 (flat_scores)",
 "0.95",
 "상위 k개 토큰 확률의 누적합(예: k=5,10). 분포가 소수 토큰에 집중됐는지(높음) 넓게 퍼졌는지(낮음) 측정. flat_scores에서 즉시 계산. 현재 미추출."],

["B. Distributional shape (분포 형태)", "entropy_at_parent (sibling)", "EAGLE3", "per-node(형제집합)",
 "float", "[0.0, log(n_sib)]", "Derived", "AVAIL", "per_node_features.py:119",
 "1.1",
 "한 노드의 형제(sibling)들 조건부 확률 분포의 entropy. 형제가 많고 확률이 균등하면 높음=경쟁 심함. 단 현재 EAGLE 계열 노드에만 채워짐(suffix 노드는 None)."],

["B. Distributional shape (분포 형태)", "suffix_sibling_sharpness (count 분포 entropy)", "suffix", "per-node(형제집합)",
 "float", "[0.0, log(n_child)]", "Derived(native count)", "PLUMB", "suffix_tree.h:55-66 (children/Group)",
 "0.3",
 "Suffix tree에서 어떤 노드의 자식들 count 분포의 sharpness/entropy. 한 자식이 지배(sharp=낮은 entropy) vs 여러 자식이 균등(모호=높은 entropy). Node.children/Group 구조에 native 존재하나 Draft엔 없음 → plumbing 필요. 단일 prob보다 '다음 토큰 결정성'을 잘 드러냄."],

["B. Distributional shape (분포 형태)", "sibling_rank (EAGLE)", "EAGLE3", "per-node",
 "int or None", "[0, n_sib)", "Derived", "AVAIL", "per_node_features.py:74",
 "0",
 "형제들 중 조건부 확률 기준 순위(0=최상위). 1등이 아니면 채택 가능성↓. 현재 EAGLE 노드만."],

["B. Distributional shape (분포 형태)", "sibling_rank (suffix, head_child 순서)", "suffix", "per-node",
 "int", "[0, n_child)", "Base", "PLUMB", "suffix_tree.h:61 (head_child)",
 "0",
 "Suffix tree는 자식을 count 내림차순 이중연결리스트로 유지(head_child=최고 count). 형제 순위가 구조적으로 native하나 Draft 미노출 → plumbing 필요."],

# ---------------------------------------------------------------- C. 빈도 / Frequency
["C. Frequency (빈도/통계)", "counts (suffix raw frequency)", "suffix", "per-node",
 "list[int]", "[1, inf)", "Base", "AVAIL", "suffix_tree.h:117 / Draft.counts",
 "[120, 80, 80, ...]",
 "Suffix tree 노드를 지나는 suffix(패턴 출현) 개수 = raw 빈도. 절대 빈도가 클수록 그 패턴이 코퍼스에서 자주 등장 → 신뢰 신호. probs의 분모/분자를 복원하는 원천값."],

["C. Frequency (빈도/통계)", "score (suffix draft score)", "suffix", "per-draft",
 "float", "[0.0, n_token]", "Native-derived", "AVAIL", "suffix_tree.h:120 / Draft.score",
 "3.4",
 "Draft 내 모든 토큰 prob의 합. ArcticInference가 직접 계산하는 draft 전체 품질/길이 가중 스칼라. local vs global tree 중 어느 draft를 쓸지 고르는 기준(cache.py)이기도 함."],

["C. Frequency (빈도/통계)", "match_len (suffix)", "suffix", "per-draft",
 "int", "[0, max_depth]", "Native-derived", "AVAIL", "suffix_tree.h:123 / Draft.match_len",
 "5",
 "Speculation 시작 전 context가 suffix tree에서 매치된 깊이(prefix 매치 길이). 길수록 더 긴 문맥이 과거 패턴과 일치 → draft 신뢰도↑. 최대 speculation 길이 제한(max_spec_factor*match_len+offset)에도 사용."],

# ---------------------------------------------------------------- D. Provenance / 출처
["D. Provenance (출처)", "tree_origin (local vs global)", "suffix", "per-draft",
 "str ('local'|'global')", "{'local','global'}", "Native(선택결과)", "PLUMB", "cache.py:489",
 "'global'",
 "Suffix cache는 request별 local tree와 누적 global tree에서 각각 draft를 만들고 score 큰 쪽을 반환. 어느 tree가 이겼는지는 native하나 현재 미기록. local 승리=요청 특이 패턴, global 승리=범용 패턴 신호."],

["D. Provenance (출처)", "tree_score_margin (local vs global)", "suffix", "per-draft",
 "float", "[0.0, inf)", "Native-derived", "PLUMB", "cache.py:489",
 "0.05",
 "두 tree draft score의 차이. margin이 작으면 어느 tree를 골라도 비슷=tree 선택 자체가 불확실. 크면 한쪽 증거가 압도적. 현재 승자만 반환하고 margin은 버려짐."],

["D. Provenance (출처)", "truncation_reason", "suffix", "per-draft",
 "str/enum", "{'confidence','leaf','length_cap'}", "Native(분기)", "PLUMB", "suffix_tree.cc (speculate 루프)",
 "'confidence'",
 "Speculation이 멈춘 이유: (a) prob<min_token_prob(confidence 한계, 고품질 가지치기), (b) 자식 없음(leaf=암기된 패턴 끝), (c) length cap 도달. '왜 멈췄나'가 draft 신뢰/특성 신호. 현재 로직엔 있으나 Draft에 미반환."],

["D. Provenance (출처)", "source", "extension", "per-node",
 "str", "{'eagle','suffix','fused'}", "Base(metadata)", "AVAIL", "tree_utils.py:25",
 "'suffix'",
 "융합 트리에서 이 노드를 만든 proposer 출처 라벨. 어떤 proposer가 낸 후보인지에 따라 다른 신호/신뢰 프로파일 적용 가능."],

# ---------------------------------------------------------------- E. 구조 / Structural
["E. Structural (구조/토폴로지)", "depth", "all", "per-node",
 "int", "[0, max_depth]", "Base", "AVAIL", "tree_utils.py:20",
 "3",
 "Root에서의 거리(트리 깊이). accept rate는 depth가 깊어질수록 감쇠하는 것이 일반적 → 위치 기반 강력한 신호. parent로부터 O(1) 계산."],

["E. Structural (구조/토폴로지)", "n_siblings", "all", "per-node",
 "int", "[0, ...)", "Base", "AVAIL", "build_per_node_dataset.py:336",
 "2",
 "같은 부모를 공유하는 다른 형제 노드 수. 형제가 많을수록 budget 경쟁/분기 불확실성↑. calibration feature(log1p)로도 사용 중."],

["E. Structural (구조/토폴로지)", "n_descendants (subtree size)", "all", "per-node",
 "int", "[0, tree_size)", "Base", "AVAIL", "per_node_features.py:34",
 "7",
 "이 노드 아래 전체 자손 수(subtree 크기). 한 노드를 선택하면 그 아래 budget이 함께 '묶여' 들어감 → budget commitment/중요도 신호. knapsack 비용으로 활용 가능."],

["E. Structural (구조/토폴로지)", "n_children (branching factor)", "all", "per-node",
 "int", "[0, ...)", "Base", "AVAIL", "per_node_features.py:25",
 "3",
 "직속 자식 수(즉각 분기 폭). 트리 내 의사결정 분기점 위치를 나타냄."],

["E. Structural (구조/토폴로지)", "node_length (path compression)", "suffix", "per-node",
 "int", "[1, ...)", "Base", "PLUMB", "suffix_tree.h:40",
 "1",
 "Suffix tree의 path-compression으로 한 노드에 묶인 토큰 수. length=1이면 토큰 단위 edge, >1이면 압축된 연속 run. 트리 granularity 신호. Draft 미노출."],

["E. Structural (구조/토폴로지)", "ref_seq (reference sequence id)", "suffix", "per-node",
 "int", "[0, n_seq)", "Base", "PLUMB", "suffix_tree.h:44",
 "12",
 "이 노드 토큰이 복사된 과거 시퀀스 ID(path compression 메타데이터). 어떤 과거 응답에서 유래했는지 추적 가능. Draft 미노출."],

# ---------------------------------------------------------------- F. Anchor (extension 전용)
["F. Anchor (extension 전용)", "ext_anchor_ids", "extension", "per-node",
 "list[int]", "{-2,-1} ∪ [0,n)", "Base", "AVAIL", "run_tree_oracle_sim.py:1522",
 "[-2, -2, 5, 5, -1]",
 "각 노드가 어디에 anchored 됐는지: -2=backbone 노드, -1=virtual root graft, >=0=graft가 붙은 backbone 노드 index. suffix graft가 backbone 트리의 어느 지점에서 뻗어나갔는지."],

["F. Anchor (extension 전용)", "anchor_depth", "extension", "per-node(graft)",
 "int", "[0, max_depth]", "Derived", "AVAIL", "build_step_dataset.py:472",
 "2",
 "Graft가 붙은 anchor 노드의 backbone 내 깊이. 깊은 anchor일수록 suffix가 더 멀리 뻗는 것."],

["F. Anchor (extension 전용)", "anchor_path_prob_eagle", "extension", "per-node(graft)",
 "float", "[0.0, 1.0]", "Derived", "AVAIL", "run_tree_oracle_sim.py:1891",
 "0.6",
 "Anchor 지점까지의 backbone(EAGLE3) 누적 경로 확률. anchor가 놓인 backbone 경로가 약하면(낮은 path_p_t) 거기 붙은 graft도 위험. graft 신뢰의 전제 조건."],

["F. Anchor (extension 전용)", "suffix_candidate_len (grafts per anchor)", "extension", "per-anchor",
 "int", "[1, ...)", "Derived", "AVAIL", "build_step_dataset.py:301",
 "4",
 "한 anchor에 매달린 suffix graft 후보 수. 인기 anchor(많은 graft)는 구조적으로 중요하거나 캐시 경합 지점일 수 있음."],

# ---------------------------------------------------------------- G. Cross-proposer
["G. Cross-proposer (proposer 간)", "proposer_agreement (same token)", "extension", "per-node",
 "bool / int", "{0,1} or count", "Derived(native topo)", "PLUMB", "run_tree_oracle_sim.py:1682",
 "True",
 "동일 (parent, token) 위치에서 두 proposer(eagle3+suffix)가 같은 토큰을 제안했는지. 둘 다 제안한 토큰은 accept 확률↑(다중 proposer 세팅 고유 신호). dedup 로직에 정보가 있으나 명시적 카운터로는 미저장."],

["G. Cross-proposer (proposer 간)", "dedup_extra", "extension", "per-node",
 "float", "[0.0, 1.0]", "Derived", "AVAIL", "run_tree_oracle_sim.py:1597",
 "0.0",
 "Backbone 노드에 겹쳐 merge된 suffix edge 확률(겹침이 없으면 0). 두 proposer가 같은 토큰에 동의할 때 suffix측 confidence를 정량화. online calibrator가 max(base, λ·dedup_extra)로 결합."],

# ---------------------------------------------------------------- H. 비용 / Cost (budget 제약)
["H. Cost (budget 제약)", "ext_tree_size", "extension", "per-step",
 "int", "[1, ...)", "Derived", "AVAIL", "run_tree_oracle_sim.py:1464",
 "48",
 "Verify해야 하는 확장 트리 전체 노드 수(backbone+graft). selection의 budget 제약(verify FLOPs/슬롯)을 직접 정의. real-cost latency 함수 입력."],

["H. Cost (budget 제약)", "subtree_cost (= n_descendants)", "all", "per-node",
 "int", "[0, ...)", "Derived", "AVAIL", "per_node_features.py:34",
 "7",
 "한 노드 선택 시 함께 verify되는 자손 수 = 그 노드의 budget 비용. knapsack에서 '가치/비용' 비율 계산용 비용 항."],

["H. Cost (budget 제약)", "verify_latency_ms", "all", "per-step",
 "float", "[0.0, ...)", "Hardware-calib", "AVAIL", "per_node_features.py:177",
 "12.3",
 "트리 크기별 target forward 실측 latency(ms). 실제 wall-clock 비용. budget을 latency 단위로 환산할 때 사용."],

# ---------------------------------------------------------------- 미존재 / 미노출 (확인 완료)
["Z. Not available (확인됨)", "recency / time-decay / node-level LRU", "suffix", "-",
 "N/A", "N/A", "-", "NONE", "cache.py (확인)",
 "-",
 "Suffix tree는 노드 단위 timestamp/LRU/시간 감쇠를 추적하지 않음. eviction은 request 단위 FIFO. → recency 신호는 존재하지 않음(추가 구현 없이는 사용 불가)."],

["Z. Not available (확인됨)", "EAGLE3 hidden-state feature (low/mid/high fusion)", "EAGLE3", "per-node",
 "tensor", "-", "-", "HIDDEN", "draft forward 내부",
 "-",
 "EAGLE3가 low/mid/high layer hidden state를 융합하지만, 이 feature는 verify 인터페이스(draft_token/path_probs/retrive_*)로 노출되지 않음. 쓰려면 draft forward를 monkey-patch해 activation을 stash해야 함."],
]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "signal_catalog"

# styles
hdr_fill = PatternFill("solid", fgColor="1F4E78")
hdr_font = Font(bold=True, color="FFFFFF", size=11)
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="top")

cat_colors = {
    "A": "DDEBF7", "B": "FCE4D6", "C": "E2EFDA", "D": "FFF2CC",
    "E": "EDEDED", "F": "DEEAF6", "G": "F8CBAD", "H": "D9D9D9", "Z": "F2DCDB",
}
avail_colors = {
    "AVAIL": "C6EFCE", "PLUMB": "FFEB9C", "HIDDEN": "FFC7CE", "NONE": "D9D9D9",
}

# header
for c, h in enumerate(HDR, 1):
    cell = ws.cell(1, c, h)
    cell.fill = hdr_fill
    cell.font = hdr_font
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    cell.border = border

# rows
for i, r in enumerate(rows, 1):
    cat, signal, proposer, gran, pytype, rng, bd, avail, ref, ex, desc = r
    cat_key = cat.split(".")[0].strip()
    vals = [i, cat, signal, proposer, gran, pytype, rng, bd, avail, ref, ex, desc]
    rownum = i + 1
    for c, v in enumerate(vals, 1):
        cell = ws.cell(rownum, c, v)
        cell.border = border
        cell.alignment = wrap if c in (2, 3, 10, 12) else center if c in (1, 4, 5, 6, 8, 9) else wrap
        if c == 2:  # category color
            cell.fill = PatternFill("solid", fgColor=cat_colors.get(cat_key, "FFFFFF"))
        if c == 9:  # availability color
            cell.fill = PatternFill("solid", fgColor=avail_colors.get(avail, "FFFFFF"))

# column widths
widths = [5, 30, 38, 12, 20, 16, 24, 16, 12, 34, 26, 75]
for c, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(c)].width = w

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(HDR))}{len(rows)+1}"

# --- legend sheet ---
ws2 = wb.create_sheet("legend")
legend = [
    ["Availability", "의미"],
    ["AVAIL", "지금 바로 사용 가능 — capture JSONL / Draft 구조체에 이미 노출됨"],
    ["PLUMB", "native 신호지만 plumbing 필요 — C++ Node / cache.py / flat_scores 내부에 존재하나 현재 추출/저장 안 함"],
    ["HIDDEN", "모델 내부에 존재하나 인터페이스 미노출 — deep monkey-patch 필요 (hidden states)"],
    ["NONE", "존재하지 않음 (코드 확인 완료) — 추가 구현 없이는 사용 불가"],
    ["", ""],
    ["Granularity", "의미"],
    ["per-node", "트리 노드(후보 토큰) 1개당 1값"],
    ["per-edge", "부모-자식 edge 1개당 1값 (보통 per-node 누적값에서 유도)"],
    ["per-draft", "한 speculate() 호출(draft) 전체당 1값 (스칼라)"],
    ["per-anchor", "extension에서 한 anchor 지점당 1값"],
    ["per-step", "디코딩 1 스텝(트리 1개) 전체당 1값"],
    ["", ""],
    ["Base/Derived", "의미"],
    ["Base", "직접 카운트/저장되는 원천값 (예: counts, token_id, depth)"],
    ["Native-derived", "라이브러리가 native하게 직접 계산하는 유도값 (예: score, probs, match_len)"],
    ["Derived", "원천값에서 우리가 trivial하게 계산 (예: edge_prob, entropy, margin)"],
    ["", ""],
    ["NOTE", "verify-time 값(target p_t)은 'inference-time 가용' 신호가 아니라 oracle upper-bound 전용이므로 본 후보 카탈로그에서 제외. BNS/DNS/topk/by_product 등 downstream selection 산출물도 제외(이건 신호가 아니라 알고리즘 출력)."],
]
for i, (a, b) in enumerate(legend, 1):
    ca = ws2.cell(i, 1, a); cb = ws2.cell(i, 2, b)
    if b == "의미" or a in ("Availability", "Granularity", "Base/Derived", "NOTE"):
        ca.font = Font(bold=True); cb.font = Font(bold=True)
    cb.alignment = Alignment(wrap_text=True, vertical="top")
ws2.column_dimensions["A"].width = 18
ws2.column_dimensions["B"].width = 95

out = "/home/muchwater/advance-spec/signal_catalog.xlsx"
wb.save(out)
print("saved:", out)
print("rows:", len(rows))
